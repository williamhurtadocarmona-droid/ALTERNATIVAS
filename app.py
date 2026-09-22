import streamlit as st
import pandas as pd
import openpyxl
from openpyxl.cell.cell import Cell
import io
import os

st.set_page_config(page_title="Generador GFPI-F-165 SENA", page_icon="📊", layout="centered")

st.title("📊 Generador de Formato GFPI-F-165 (Etapa Productiva)")
st.write("Sube el reporte de Sofia Plus para llenar el formato **Grupal** y generar las pestañas **Individuales**.")

# Función para escribir de forma segura en celdas combinadas
def escribir_celda_segura(hoja, coordenada, valor):
    celda = hoja[coordenada]
    if isinstance(celda, Cell):
        celda.value = valor
    else:
        # Si la celda es parte de un rango combinado, buscamos la celda principal (top-left)
        for rango in hoja.merged_cells.ranges:
            if coordenada in rango:
                hoja.cell(row=rango.min_row, column=rango.min_col, value=valor)
                break

# 1. Cargar el reporte de Sofia Plus
excel_reporte = st.file_uploader("Sube el reporte de Sofia Plus (.xlsx / .xls)", type=["xlsx", "xls"])

PLANTILLA_BASE = "GFPI-F-165.xlsx"  # Archivo base en GitHub

if excel_reporte is not None:
    try:
        if not os.path.exists(PLANTILLA_BASE):
            st.error(f"⚠️ No se encontró la plantilla base '{PLANTILLA_BASE}' en GitHub. Asegúrate de que el archivo tenga ese nombre exacto.")
        else:
            # 2. Leer Encabezado del reporte Sofia Plus (filas 1-12)
            df_encabezado = pd.read_excel(excel_reporte, nrows=12, header=None)
            
            centro = str(df_encabezado.iloc[1, 2]).strip() if len(df_encabezado) > 1 and pd.notna(df_encabezado.iloc[1, 2]) else ""
            ficha = str(df_encabezado.iloc[2, 2]).strip() if len(df_encabezado) > 2 and pd.notna(df_encabezado.iloc[2, 2]) else ""
            programa = str(df_encabezado.iloc[5, 2]).strip() if len(df_encabezado) > 5 and pd.notna(df_encabezado.iloc[5, 2]) else ""

            # 3. Leer la tabla de aprendices
            df_aprendices = pd.read_excel(excel_reporte, skiprows=12)
            df_aprendices.columns = [str(c).strip() for c in df_aprendices.columns]

            # Detectar columnas automáticamente
            col_estado = [c for c in df_aprendices.columns if 'Estado' in c or 'ESTADO' in c]
            col_estado = col_estado[0] if col_estado else df_aprendices.columns[-1]

            col_num_doc = [c for c in df_aprendices.columns if ('Número' in c or 'Numero' in c or 'Documento' in c or 'Identificación' in c) and 'Tipo' not in c]
            col_num_doc = col_num_doc[0] if col_num_doc else None

            col_nombres = [c for c in df_aprendices.columns if 'Nombre' in c or 'Aprendiz' in c]
            col_apellidos = [c for c in df_aprendices.columns if 'Apellido' in c]
            
            col_tipo_doc = [c for c in df_aprendices.columns if 'Tipo' in c]
            col_tipo_doc = col_tipo_doc[0] if col_tipo_doc else None

            col_correo = [c for c in df_aprendices.columns if 'Correo' in c or 'Email' in c]
            col_correo = col_correo[0] if col_correo else None

            col_tel = [c for c in df_aprendices.columns if 'Teléfono' in c or 'Celular' in c or 'Móvil' in c]
            col_tel = col_tel[0] if col_tel else None

            # 4. Filtrar aprendices activos (EN FORMACIÓN)
            df_activos = df_aprendices[df_aprendices[col_estado].astype(str).str.upper().str.contains("EN FORMAC", na=False)].copy()

            col_clave = col_num_doc if col_num_doc else (col_nombres[0] if col_nombres else None)

            if col_clave:
                df_unicos = df_activos.drop_duplicates(subset=[col_clave]).copy()
            else:
                df_unicos = df_activos.copy()

            cant_totales_filas = len(df_aprendices)
            cant_unicos_activos = len(df_unicos)

            st.info(f"📋 **Resumen:** Filas en reporte: {cant_totales_filas} | Aprendices ÚNICOS en **EN FORMACIÓN**: **{cant_unicos_activos}**")

            if cant_unicos_activos == 0:
                st.warning("⚠️ No se encontraron aprendices con estado 'EN FORMACIÓN' en el archivo.")
            else:
                cols_mostrar = [c for c in [col_tipo_doc, col_num_doc, col_nombres[0] if col_nombres else None, col_estado] if c is not None]
                st.dataframe(df_unicos[cols_mostrar])

                if st.button("🚀 Generar Formato Grupal y Pestañas Individuales"):
                    with st.spinner(f"Procesando {cant_unicos_activos} aprendices..."):
                        
                        wb = openpyxl.load_workbook(PLANTILLA_BASE)
                        
                        # ==========================================
                        # A. LLENAR PESTAÑA GRUPAL
                        # ==========================================
                        HOJA_GRUPAL_NOMBRE = "Selección formato 1 - Grupal"
                        if HOJA_GRUPAL_NOMBRE in wb.sheetnames:
                            hoja_grupal = wb[HOJA_GRUPAL_NOMBRE]
                            
                            fila_inicio_grupal = 13 
                            
                            for i, (_, row) in enumerate(df_unicos.iterrows(), start=0):
                                fila_actual = fila_inicio_grupal + i
                                
                                nomb = str(row[col_nombres[0]]).strip() if col_nombres and pd.notna(row[col_nombres[0]]) else ""
                                apel = str(row[col_apellidos[0]]).strip() if col_apellidos and pd.notna(row[col_apellidos[0]]) else ""
                                tipo_doc = str(row[col_tipo_doc]).strip() if col_tipo_doc and pd.notna(row[col_tipo_doc]) else ""
                                num_doc = str(row[col_num_doc]).strip() if col_num_doc and pd.notna(row[col_num_doc]) else ""

                                escribir_celda_segura(hoja_grupal, f'A{fila_actual}', i + 1)
                                escribir_celda_segura(hoja_grupal, f'B{fila_actual}', tipo_doc)
                                escribir_celda_segura(hoja_grupal, f'C{fila_actual}', num_doc)
                                escribir_celda_segura(hoja_grupal, f'D{fila_actual}', nomb)
                                escribir_celda_segura(hoja_grupal, f'E{fila_actual}', apel)
                                
                                if col_correo and pd.notna(row[col_correo]):
                                    escribir_celda_segura(hoja_grupal, f'G{fila_actual}', str(row[col_correo]))
                                if col_tel and pd.notna(row[col_tel]):
                                    escribir_celda_segura(hoja_grupal, f'H{fila_actual}', str(row[col_tel]))

                        # ==========================================
                        # B. GENERAR PESTAÑAS INDIVIDUALES
                        # ==========================================
                        NOMBRE_HOJA_PLANTILLA = "Selección Modificación F2 Indiv"

                        if NOMBRE_HOJA_PLANTILLA in wb.sheetnames:
                            hoja_base = wb[NOMBRE_HOJA_PLANTILLA]
                        else:
                            hoja_encontrada = [s for s in wb.sheetnames if "F2" in s or "Indiv" in s]
                            if hoja_encontrada:
                                hoja_base = wb[hoja_encontrada[0]]
                            else:
                                st.error(f"⚠️ No se encontró la pestaña '{NOMBRE_HOJA_PLANTILLA}' en la plantilla.")
                                st.stop()

                        for idx, row in df_unicos.iterrows():
                            nomb = str(row[col_nombres[0]]).strip() if col_nombres and pd.notna(row[col_nombres[0]]) else ""
                            apel = str(row[col_apellidos[0]]).strip() if col_apellidos and pd.notna(row[col_apellidos[0]]) else ""
                            nombre_completo = f"{nomb} {apel}".strip()

                            tipo_doc = str(row[col_tipo_doc]).strip() if col_tipo_doc and pd.notna(row[col_tipo_doc]) else ""
                            num_doc = str(row[col_num_doc]).strip() if col_num_doc and pd.notna(row[col_num_doc]) else ""

                            id_pestana = num_doc[-4:] if num_doc else nomb.split()[0]
                            titulo_pestaña = f"{nomb.split()[0]} {id_pestana}"[:30]

                            target_sheet = wb.copy_worksheet(hoja_base)
                            target_sheet.title = titulo_pestaña

                            escribir_celda_segura(target_sheet, 'C12', tipo_doc)
                            escribir_celda_segura(target_sheet, 'D12', num_doc)
                            escribir_celda_segura(target_sheet, 'E12', nombre_completo)
                            
                            if col_tel and pd.notna(row[col_tel]):
                                escribir_celda_segura(target_sheet, 'F12', str(row[col_tel]))
                            if col_correo and pd.notna(row[col_correo]):
                                escribir_celda_segura(target_sheet, 'G12', str(row[col_correo]))
                            
                            escribir_celda_segura(target_sheet, 'C17', centro)
                            escribir_celda_segura(target_sheet, 'E17', ficha)
                            escribir_celda_segura(target_sheet, 'F17', programa)

                        # Guardar resultado
                        output = io.BytesIO()
                        wb.save(output)
                        output.seek(0)

                        st.download_button(
                            label="📥 Descargar Libro Excel Consolidado (.xlsx)",
                            data=output,
                            file_name=f"GFPI-F-165_Ficha_{ficha}_Consolidado.xlsx",
                            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                        )
                        st.success(f"¡Listo! Se procesó el formato Grupal y se generaron {cant_unicos_activos} pestañas individuales sin errores.")

    except Exception as e:
        st.error(f"Ocurrió un error al procesar el archivo: {e}")
